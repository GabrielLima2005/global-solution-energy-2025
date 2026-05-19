# -*- coding: utf-8 -*-
"""
Questão 4 - Roteiro do Vídeo Pitch (~1 min)
PCP - Sprint 3
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

def set_font(run, name='Times New Roman', size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)
    rPr.insert(0, rFonts)

def add_centered_text(text, size=12, bold=False, space_after=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = Pt(18)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold)
    return p

def add_normal_text(text, space_after=6, bold=False, indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = Pt(18)
    if indent:
        pf.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_font(run, size=12, bold=bold)
    return p

def add_heading_text(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_font(run, size=14 if level == 1 else 12, bold=True)
    return h

def add_script_block(tempo, titulo, fala, dica_visual=''):
    """Adiciona bloco de roteiro com tempo, título, fala e dica visual"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(2)
    pf.line_spacing = Pt(18)

    run = p.add_run(f'[{tempo}] ')
    set_font(run, size=11, bold=True, color=(0, 102, 204))

    run = p.add_run(titulo)
    set_font(run, size=11, bold=True)

    # Fala
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf2 = p2.paragraph_format
    pf2.space_after = Pt(2)
    pf2.line_spacing = Pt(18)
    pf2.left_indent = Cm(1)

    run = p2.add_run('\u201c')
    set_font(run, size=11, italic=True)
    run = p2.add_run(fala)
    set_font(run, size=11, italic=True)
    run = p2.add_run('\u201d')
    set_font(run, size=11, italic=True)

    # Dica visual
    if dica_visual:
        p3 = doc.add_paragraph()
        pf3 = p3.paragraph_format
        pf3.space_after = Pt(6)
        pf3.line_spacing = Pt(16)
        pf3.left_indent = Cm(1)
        run = p3.add_run('[Visual] ')
        set_font(run, size=10, bold=True, color=(128, 128, 128))
        run = p3.add_run(dica_visual)
        set_font(run, size=10, color=(128, 128, 128))

# =================== CAPA ===================
add_centered_text('FIAP \u2013 FACULDADE DE INFORM\u00c1TICA E ADMINISTRA\u00c7\u00c3O PAULISTA', 14, bold=True, space_after=24)
add_centered_text('1\u00ba ANO DE CI\u00caNCIA DA COMPUTA\u00c7\u00c3O', 12, space_after=12)
add_centered_text('PCP \u2013 PENSAMENTO COMPUTACIONAL E AUTOMA\u00c7\u00c3O COM PYTHON', 12, space_after=12)
add_centered_text('Prof. Allan Roberto Molto', 12, space_after=48)

add_centered_text('QUEST\u00c3O 4 \u2013 ROTEIRO DO V\u00cdDEO PITCH', 16, bold=True, space_after=6)
add_centered_text('Level UP \u2013 BluaDiagnostics', 14, bold=True, space_after=6)
add_centered_text('Dura\u00e7\u00e3o: ~1 minuto', 12, space_after=48)

add_centered_text('Nicolas Ara\u00fajo de Oliveira \u2013 RM 566780', 12, space_after=6)
add_centered_text('Pedro Ivson Falc\u00e3o De Leucas', 12, space_after=6)
add_centered_text('Gabriel Lima Da Silva \u2013 RM 568436', 12, space_after=48)

add_centered_text('S\u00e3o Paulo', 12, space_after=6)
add_centered_text('2025', 12)

doc.add_page_break()

# =================== ROTEIRO ===================
add_heading_text('ROTEIRO DO V\u00cdDEO PITCH (~1 MINUTO)', level=1)

add_normal_text(
    'Abaixo est\u00e1 o roteiro completo do v\u00eddeo pitch, dividido em se\u00e7\u00f5es com '
    'marca\u00e7\u00e3o de tempo, texto da fala e dicas visuais para a grava\u00e7\u00e3o.',
    indent=True
)

p = doc.add_paragraph()
pf = p.paragraph_format
pf.space_after = Pt(4)
run = p.add_run('Tempo total estimado: 55\u201365 segundos')
set_font(run, size=11, bold=True)

p = doc.add_paragraph()
pf = p.paragraph_format
pf.space_after = Pt(12)
run = p.add_run('Dica: fale com ritmo firme e entusiasmo controlado. Pratique antes de gravar!')
set_font(run, size=10, italic=True, color=(128, 128, 128))

# --- ABERTURA ---
add_script_block(
    '0:00 \u2013 0:08',
    'ABERTURA \u2013 O Problema',
    'Hoje, 60% das doen\u00e7as cr\u00f4nicas poderiam ser evitadas com monitoramento adequado. '
    'Mas o modelo de sa\u00fade atual ainda \u00e9 reativo: o paciente s\u00f3 busca atendimento quando '
    'j\u00e1 est\u00e1 doente. E se pud\u00e9ssemos mudar isso?',
    'Tela preta com texto "60% das doen\u00e7as cr\u00f4nicas s\u00e3o evit\u00e1veis" aparecendo gradualmente'
)

# --- SOLUCAO ---
add_script_block(
    '0:08 \u2013 0:15',
    'A SOLU\u00c7\u00c3O \u2013 Level UP',
    'Apresentamos o Level UP: uma evolu\u00e7\u00e3o da plataforma Blua da CarePlus que transforma '
    'o cuidado de sa\u00fade em um modelo inteligente, proativo e centrado no paciente.',
    'Logo do Level UP aparece. Transi\u00e7\u00e3o para o circuito no Wokwi'
)

# --- PILAR 1 ---
add_script_block(
    '0:15 \u2013 0:25',
    'PILAR 1 \u2013 Monitoramento Cont\u00ednuo',
    'O primeiro pilar \u00e9 o monitoramento cont\u00ednuo via wearables. Sensores coletam '
    'frequ\u00eancia card\u00edaca, temperatura e atividade f\u00edsica em tempo real. LEDs indicam '
    'o status: verde para normal, amarelo para aten\u00e7\u00e3o e vermelho com alarme para '
    'situa\u00e7\u00f5es cr\u00edticas.',
    'Mostrar o circuito com LEDs acendendo. Mudar potenci\u00f4metro para mostrar LED vermelho + buzzer'
)

# --- PILAR 2 ---
add_script_block(
    '0:25 \u2013 0:35',
    'PILAR 2 \u2013 Check-up Digital com IA',
    'O segundo pilar \u00e9 o check-up digital com intelig\u00eancia artificial. O paciente responde '
    '8 perguntas de triagem no display. O sistema cruza as respostas com os sinais vitais e '
    'gera um score de risco de 0 a 100, com classifica\u00e7\u00e3o autom\u00e1tica e relat\u00f3rio m\u00e9dico.',
    'Mostrar tela do check-up no OLED. Responder perguntas. Exibir resultado do score'
)

# --- PILAR 3 ---
add_script_block(
    '0:35 \u2013 0:42',
    'PILAR 3 \u2013 Gamifica\u00e7\u00e3o',
    'E o terceiro pilar \u00e9 a gamifica\u00e7\u00e3o! O paciente ganha pontos por manter h\u00e1bitos '
    'saud\u00e1veis, sobe de n\u00edvel e conquista ranks como "Guerreiro Fit" e "Mestre Sa\u00fade". '
    'Isso aumenta a ades\u00e3o ao tratamento em at\u00e9 30%.',
    'Mostrar tela de gamifica\u00e7\u00e3o no OLED com pontos subindo'
)

# --- DASHBOARD ---
add_script_block(
    '0:42 \u2013 0:50',
    'DASHBOARD WEB',
    'Al\u00e9m do display OLED, o ESP32 cria sua pr\u00f3pria rede WiFi e serve um dashboard web '
    'responsivo com todos os dados em tempo real. Basta conectar o celular e acessar o '
    'painel pelo navegador.',
    'Mostrar o dashboard web no celular/navegador com os cards atualizando'
)

# --- ENCERRAMENTO ---
add_script_block(
    '0:50 \u2013 1:00',
    'ENCERRAMENTO \u2013 Impacto',
    'O Level UP transforma o cuidado de reativo para proativo. Para a CarePlus, isso '
    'significa redu\u00e7\u00e3o de custos assistenciais, aumento da satisfa\u00e7\u00e3o dos benefici\u00e1rios '
    'e fortalecimento como refer\u00eancia em inova\u00e7\u00e3o em sa\u00fade. Level UP: sua sa\u00fade, '
    'seu jogo, seu n\u00edvel!',
    'Tela final com logo Level UP e slogan. Nomes dos integrantes da equipe'
)

doc.add_page_break()

# =================== DICAS DE GRAVACAO ===================
add_heading_text('DICAS PARA A GRAVA\u00c7\u00c3O', level=1)

add_normal_text('1. Equipamento:', bold=True)
add_normal_text(
    '\u2022 Use um celular com boa c\u00e2mera em modo horizontal (paisagem)\n'
    '\u2022 Grave em ambiente silencioso e bem iluminado\n'
    '\u2022 Use trip\u00e9 ou apoie o celular para evitar tremores'
)

add_normal_text('2. Apresenta\u00e7\u00e3o:', bold=True)
add_normal_text(
    '\u2022 Fale olhando para a c\u00e2mera, com ritmo firme\n'
    '\u2022 Evite ler o roteiro \u2013 pratique antes para soar natural\n'
    '\u2022 Sorria e demonstre entusiasmo pelo projeto'
)

add_normal_text('3. Demonstra\u00e7\u00e3o:', bold=True)
add_normal_text(
    '\u2022 Tenha o Wokwi aberto no computador para demonstrar o circuito\n'
    '\u2022 Prepare o dashboard web j\u00e1 aberto no navegador\n'
    '\u2022 Mostre os LEDs mudando de cor e o buzzer ativando\n'
    '\u2022 Realize um check-up completo durante a demonstra\u00e7\u00e3o'
)

add_normal_text('4. Edi\u00e7\u00e3o:', bold=True)
add_normal_text(
    '\u2022 Adicione t\u00edtulos e transi\u00e7\u00f5es entre as se\u00e7\u00f5es\n'
    '\u2022 Inclua legendas se poss\u00edvel (acessibilidade)\n'
    '\u2022 Mantenha o v\u00eddeo entre 55 e 65 segundos\n'
    '\u2022 Adicione m\u00fasica de fundo suave (royalty-free)'
)

add_normal_text('5. Entrega:', bold=True)
add_normal_text(
    '\u2022 Exporte em MP4 (1080p recomendado)\n'
    '\u2022 Nomeie o arquivo: Video_Pitch_LevelUP.mp4\n'
    '\u2022 Verifique o \u00e1udio antes de enviar'
)

doc.add_page_break()

# =================== STORYBOARD ===================
add_heading_text('STORYBOARD VISUAL', level=1)

add_normal_text(
    'Abaixo, o storyboard simplificado com as cenas principais do v\u00eddeo:',
    indent=True
)

from docx.enum.table import WD_TABLE_ALIGNMENT

table = doc.add_table(rows=7, cols=3)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ['Cena', 'Tempo', 'Descri\u00e7\u00e3o Visual']
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(h)
    set_font(run, size=10, bold=True)
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), '1a1a3e')
    shading.set(qn('w:val'), 'clear')
    cell._element.get_or_add_tcPr().append(shading)
    run.font.color.rgb = RGBColor(255, 255, 255)

scenes = [
    ['1', '0:00\u20130:08', 'Tela escura com estat\u00edstica impactante. Apresentador entra em cena questionando o modelo atual de sa\u00fade.'],
    ['2', '0:08\u20130:15', 'Logo Level UP. Circuito Wokwi aparece em tela dividida com o apresentador.'],
    ['3', '0:15\u20130:25', 'Close no circuito: LEDs mudando de cor (verde \u2192 amarelo \u2192 vermelho). Potenci\u00f4metros sendo girados. Buzzer ativa.'],
    ['4', '0:25\u20130:35', 'Tela do OLED mostrando check-up. Bot\u00f5es sendo pressionados. Resultado aparece com score e classifica\u00e7\u00e3o.'],
    ['5', '0:35\u20130:42', 'Tela de gamifica\u00e7\u00e3o no OLED. Anima\u00e7\u00e3o de "LEVEL UP!" com estrelas. Barra de progresso enchendo.'],
    ['6', '0:42\u20131:00', 'Dashboard web no celular/notebook. Cards atualizando em tempo real. Encerramento com slogan e cr\u00e9ditos.'],
]

for r_idx, scene in enumerate(scenes):
    for c_idx, val in enumerate(scene):
        cell = table.rows[r_idx + 1].cells[c_idx]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 2 else WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(val)
        set_font(run, size=10)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pf = p.paragraph_format
pf.space_before = Pt(4)
pf.space_after = Pt(12)
run = p.add_run('Quadro 1 \u2013 Storyboard do v\u00eddeo pitch')
set_font(run, size=10, italic=True)

# =================== SALVAR ===================
output_path = '/home/ubuntu/repos/global-solution-energy-2025/level-up-wokwi/Questao4_Roteiro_Pitch.docx'
doc.save(output_path)
print(f'Documento salvo em: {output_path}')
print(f'Tamanho: {os.path.getsize(output_path)} bytes')
