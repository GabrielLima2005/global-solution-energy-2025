# -*- coding: utf-8 -*-
"""
Questão 2 - Cronograma detalhado em Excel
PCP - Sprint 3
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os

wb = Workbook()
ws = wb.active
ws.title = 'Cronograma Level UP'

# =================== ESTILOS ===================
header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='1a1a3e', end_color='1a1a3e', fill_type='solid')
subheader_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
subheader_fill = PatternFill(start_color='2d4a7a', end_color='2d4a7a', fill_type='solid')
data_font = Font(name='Calibri', size=10)
done_fill = PatternFill(start_color='c6efce', end_color='c6efce', fill_type='solid')
progress_fill = PatternFill(start_color='ffeb9c', end_color='ffeb9c', fill_type='solid')
future_fill = PatternFill(start_color='bdd7ee', end_color='bdd7ee', fill_type='solid')
status_done_font = Font(name='Calibri', size=10, bold=True, color='006100')
status_progress_font = Font(name='Calibri', size=10, bold=True, color='9c6500')
status_future_font = Font(name='Calibri', size=10, bold=True, color='1f4e79')
center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# =================== TITULO ===================
ws.merge_cells('A1:G1')
title_cell = ws['A1']
title_cell.value = 'CRONOGRAMA DETALHADO \u2013 PROJETO LEVEL UP (Care Plus / BluaDiagnostics)'
title_cell.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
title_cell.fill = PatternFill(start_color='0d0d2b', end_color='0d0d2b', fill_type='solid')
title_cell.alignment = center_align

ws.merge_cells('A2:G2')
subtitle_cell = ws['A2']
subtitle_cell.value = 'PCP \u2013 Pensamento Computacional e Automa\u00e7\u00e3o com Python | 2\u00ba Semestre \u2013 Sprint 3 | Prof. Allan Roberto Molto'
subtitle_cell.font = Font(name='Calibri', size=10, italic=True, color='FFFFFF')
subtitle_cell.fill = PatternFill(start_color='1a1a3e', end_color='1a1a3e', fill_type='solid')
subtitle_cell.alignment = center_align

ws.merge_cells('A3:G3')
team_cell = ws['A3']
team_cell.value = 'Equipe: Nicolas Ara\u00fajo (RM 566780) | Pedro Ivson Falc\u00e3o De Leucas | Gabriel Lima (RM 568436)'
team_cell.font = Font(name='Calibri', size=10, color='FFFFFF')
team_cell.fill = PatternFill(start_color='1a1a3e', end_color='1a1a3e', fill_type='solid')
team_cell.alignment = center_align

# =================== CABECALHO ===================
headers = ['#', 'Etapa', 'Descri\u00e7\u00e3o Detalhada', 'Respons\u00e1vel', 'In\u00edcio', 'T\u00e9rmino', 'Status']
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=5, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border

# =================== DADOS ===================
etapas = [
    # Sprint 1 - Concluidas
    ['SPRINT 1 \u2013 PROTOTIPAGEM E CIRCUITO (CONCLU\u00cdDA)', None, None, None, None, None],

    ['1', 'Planejamento do projeto',
     'Defini\u00e7\u00e3o dos pilares (monitoramento cont\u00ednuo, check-up digital, gamifica\u00e7\u00e3o), '
     'escolha de componentes e plataforma Wokwi.',
     'Equipe', '01/04/2025', '03/04/2025', 'Conclu\u00eddo \u2713'],

    ['2', 'Desenvolvimento do circuito (diagram.json)',
     'Montagem do circuito no Wokwi: ESP32, OLED SSD1306, DHT22, 2 potenci\u00f4metros, '
     '3 LEDs (verde/amarelo/vermelho), buzzer, 2 bot\u00f5es.',
     'Gabriel', '04/04/2025', '06/04/2025', 'Conclu\u00eddo \u2713'],

    ['3', 'C\u00f3digo Arduino v1.0 (sketch.ino)',
     'Programa\u00e7\u00e3o do ESP32: leitura de sensores, 5 telas OLED (Dashboard, Vitais, '
     'Check-up, Gamifica\u00e7\u00e3o, Alerta), sistema de pontua\u00e7\u00e3o e ranks.',
     'Nicolas', '04/04/2025', '08/04/2025', 'Conclu\u00eddo \u2713'],

    ['4', 'Sistema de Check-up Digital',
     '8 perguntas ponderadas com pesos cl\u00ednicos, c\u00e1lculo de score de risco (0\u2013100), '
     'classifica\u00e7\u00e3o autom\u00e1tica e relat\u00f3rio m\u00e9dico via Serial.',
     'Pedro', '06/04/2025', '09/04/2025', 'Conclu\u00eddo \u2713'],

    ['5', 'Sistema de Gamifica\u00e7\u00e3o',
     'Pontua\u00e7\u00e3o por h\u00e1bitos saud\u00e1veis a cada 30s, 5 ranks progressivos, '
     'b\u00f4nus por check-ups completados.',
     'Gabriel', '06/04/2025', '09/04/2025', 'Conclu\u00eddo \u2713'],

    ['6', 'Testes e valida\u00e7\u00e3o v1.0',
     'Testes de todos os sensores, valida\u00e7\u00e3o do check-up, verifica\u00e7\u00e3o dos LEDs/buzzer, '
     'navega\u00e7\u00e3o entre telas.',
     'Equipe', '09/04/2025', '10/04/2025', 'Conclu\u00eddo \u2713'],

    # Sprint 2 - Concluidas
    ['SPRINT 2 \u2013 WIFI, DASHBOARD WEB E REL\u00d3GIO (CONCLU\u00cdDA)', None, None, None, None, None],

    ['7', 'Implementa\u00e7\u00e3o WiFi Access Point',
     'Configura\u00e7\u00e3o do ESP32 como Access Point (SSID: LevelUP-Health, senha: levelup123). '
     'Cria\u00e7\u00e3o de rede WiFi independente.',
     'Nicolas', '14/04/2025', '15/04/2025', 'Conclu\u00eddo \u2713'],

    ['8', 'Web Server + Dashboard HTML',
     'Servidor HTTP na porta 80 com p\u00e1gina HTML responsiva em PROGMEM. '
     'Dashboard com tema escuro, cards coloridos, barras de progresso.',
     'Gabriel', '15/04/2025', '18/04/2025', 'Conclu\u00eddo \u2713'],

    ['9', 'API REST',
     'Endpoint GET /api/data retornando JSON com todos os dados de sa\u00fade. '
     'Endpoint GET /api/settime para ajuste remoto do rel\u00f3gio.',
     'Pedro', '16/04/2025', '18/04/2025', 'Conclu\u00eddo \u2713'],

    ['10', 'Rel\u00f3gio em Tempo Real',
     'Software clock via millis() com precis\u00e3o de 1 segundo. '
     'Nova tela OLED dedicada (HH:MM SS) + exibi\u00e7\u00e3o no dashboard web.',
     'Nicolas', '18/04/2025', '19/04/2025', 'Conclu\u00eddo \u2713'],

    ['11', 'Atualiza\u00e7\u00e3o sketch.ino v2.0',
     'Integra\u00e7\u00e3o de WiFi, WebServer, rel\u00f3gio e 6\u00aa tela. '
     'C\u00f3digo final com 977 linhas.',
     'Gabriel', '19/04/2025', '20/04/2025', 'Conclu\u00eddo \u2713'],

    ['12', 'Atualiza\u00e7\u00e3o README.md',
     'Documenta\u00e7\u00e3o completa: instru\u00e7\u00f5es WiFi, API REST com exemplo JSON, '
     'guia de 6 telas OLED, tabela de componentes.',
     'Equipe', '20/04/2025', '20/04/2025', 'Conclu\u00eddo \u2713'],

    # Sprint 3 - Atual
    ['SPRINT 3 \u2013 L\u00d3GICA DIGITAL E DOCUMENTA\u00c7\u00c3O (ATUAL)', None, None, None, None, None],

    ['13', 'Modelagem da l\u00f3gica digital',
     'Identifica\u00e7\u00e3o de 3 subsistemas l\u00f3gicos: Alertas (4 vari\u00e1veis, 4 sa\u00eddas), '
     'Gamifica\u00e7\u00e3o (5 vari\u00e1veis), Classifica\u00e7\u00e3o de Risco (2 vari\u00e1veis, 3 sa\u00eddas).',
     'Nicolas', '05/05/2025', '07/05/2025', 'Conclu\u00eddo \u2713'],

    ['14', 'Simplifica\u00e7\u00e3o com Mapas de Karnaugh',
     'Aplica\u00e7\u00e3o de Karnaugh, \u00e1lgebra booleana e Teorema de De Morgan. '
     'Redu\u00e7\u00e3o de 8 para 5 portas l\u00f3gicas (-37,5%) no subsistema de alertas.',
     'Pedro', '07/05/2025', '09/05/2025', 'Conclu\u00eddo \u2713'],

    ['15', 'Tabelas verdade comparativas',
     'Elabora\u00e7\u00e3o de tabelas verdade para modelos original e simplificado. '
     'Prova de equival\u00eancia em todos os cen\u00e1rios poss\u00edveis.',
     'Gabriel', '09/05/2025', '10/05/2025', 'Conclu\u00eddo \u2713'],

    ['16', 'Relat\u00f3rio de L\u00f3gica Digital (Markdown)',
     'Documento completo com express\u00f5es booleanas, Karnaugh, tabelas verdade, '
     'diagramas l\u00f3gicos e prova de equival\u00eancia.',
     'Equipe', '10/05/2025', '11/05/2025', 'Conclu\u00eddo \u2713'],

    ['17', 'Relat\u00f3rio ABNT em Word (.docx)',
     'Formata\u00e7\u00e3o ABNT (NBR 14724): Times New Roman 12pt, espa\u00e7amento 1,5, '
     'margens 3cm/2cm, capa, sum\u00e1rio, refer\u00eancias NBR 6023.',
     'Gabriel', '11/05/2025', '12/05/2025', 'Conclu\u00eddo \u2713'],

    ['18', 'Quest\u00e3o 1 \u2013 Descri\u00e7\u00e3o Completa do Projeto',
     'Documento Word com descri\u00e7\u00e3o de dados de entrada (wearables/IA), '
     'an\u00e1lise dos dados e sa\u00eddas com benef\u00edcios.',
     'Nicolas', '14/05/2025', '16/05/2025', 'Conclu\u00eddo \u2713'],

    ['19', 'Quest\u00e3o 2 \u2013 Cronograma em Excel',
     'Planilha detalhada com todas as etapas, respons\u00e1veis, '
     'prazos e status (realizadas e futuras).',
     'Pedro', '14/05/2025', '16/05/2025', 'Conclu\u00eddo \u2713'],

    ['20', 'Quest\u00e3o 3 \u2013 Apresenta\u00e7\u00e3o com Prints',
     'Documento com prints/capturas de todas as telas do sistema '
     '(OLED, Dashboard Web, Serial, C\u00f3digo).',
     'Gabriel', '16/05/2025', '18/05/2025', 'Conclu\u00eddo \u2713'],

    ['21', 'Quest\u00e3o 4 \u2013 V\u00eddeo Pitch (~1 min)',
     'Grava\u00e7\u00e3o do v\u00eddeo pitch apresentando o projeto completo: '
     'problema, solu\u00e7\u00e3o, pilares, demonstra\u00e7\u00e3o e resultados.',
     'Equipe', '16/05/2025', '18/05/2025', 'Em andamento'],

    # Etapas Futuras
    ['ETAPAS FUTURAS \u2013 EVOLU\u00c7\u00c3O DO PROJETO', None, None, None, None, None],

    ['22', 'Integra\u00e7\u00e3o com wearables reais',
     'Substitui\u00e7\u00e3o dos potenci\u00f4metros por sensores reais: '
     'MAX30102 (ox\u00edmetro/BPM), MPU6050 (aceler\u00f4metro).',
     'Equipe', '01/06/2025', '15/06/2025', 'Planejado'],

    ['23', 'Backend em nuvem (API + Banco de Dados)',
     'Migra\u00e7\u00e3o do Web Server local para API em nuvem (Firebase/AWS). '
     'Banco de dados para hist\u00f3rico de sa\u00fade do paciente.',
     'Equipe', '15/06/2025', '30/06/2025', 'Planejado'],

    ['24', 'IA real para triagem',
     'Substituir pesos fixos por modelo de Machine Learning '
     'treinado com dados cl\u00ednicos reais para triagem inteligente.',
     'Equipe', '01/07/2025', '20/07/2025', 'Planejado'],

    ['25', 'Integra\u00e7\u00e3o com plataforma Blua',
     'Conex\u00e3o com API da Blua/CarePlus para sincroniza\u00e7\u00e3o de dados, '
     'agendamento de teleconsultas e prescri\u00e7\u00e3o digital.',
     'Equipe', '20/07/2025', '10/08/2025', 'Planejado'],

    ['26', 'Testes com usu\u00e1rios e valida\u00e7\u00e3o cl\u00ednica',
     'Testes de usabilidade com pacientes piloto, valida\u00e7\u00e3o dos '
     'algoritmos de risco com profissionais de sa\u00fade.',
     'Equipe', '10/08/2025', '30/08/2025', 'Planejado'],
]

row = 6
for etapa in etapas:
    if etapa[1] is None:
        # Section header
        ws.merge_cells(f'A{row}:G{row}')
        cell = ws.cell(row=row, column=1, value=etapa[0])
        cell.font = subheader_font
        cell.fill = subheader_fill
        cell.alignment = center_align
        cell.border = thin_border
        for c in range(2, 8):
            ws.cell(row=row, column=c).border = thin_border
        row += 1
        continue

    values = etapa
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.font = data_font
        cell.alignment = left_align if col in [2, 3] else center_align
        cell.border = thin_border

        # Color by status
        status = values[6] if len(values) > 6 else ''
        if 'Conclu\u00eddo' in str(status):
            if col == 7:
                cell.font = status_done_font
            cell.fill = done_fill
        elif 'andamento' in str(status).lower():
            if col == 7:
                cell.font = status_progress_font
            cell.fill = progress_fill
        elif 'Planejado' in str(status):
            if col == 7:
                cell.font = status_future_font
            cell.fill = future_fill

    row += 1

# =================== LARGURA DAS COLUNAS ===================
ws.column_dimensions['A'].width = 5
ws.column_dimensions['B'].width = 32
ws.column_dimensions['C'].width = 55
ws.column_dimensions['D'].width = 14
ws.column_dimensions['E'].width = 14
ws.column_dimensions['F'].width = 14
ws.column_dimensions['G'].width = 16

# =================== LEGENDA ===================
row += 1
ws.merge_cells(f'A{row}:G{row}')
legend = ws.cell(row=row, column=1, value='LEGENDA:')
legend.font = Font(name='Calibri', size=10, bold=True)

row += 1
for c in range(1, 4):
    ws.cell(row=row, column=c).border = thin_border

ws.cell(row=row, column=1, value='\u2713 Conclu\u00eddo').fill = done_fill
ws.cell(row=row, column=1).font = status_done_font
ws.cell(row=row, column=2, value='\u25cb Em andamento').fill = progress_fill
ws.cell(row=row, column=2).font = status_progress_font
ws.cell(row=row, column=3, value='\u25a1 Planejado').fill = future_fill
ws.cell(row=row, column=3).font = status_future_font

# =================== SALVAR ===================
output_path = '/home/ubuntu/repos/global-solution-energy-2025/level-up-wokwi/Questao2_Cronograma.xlsx'
wb.save(output_path)
print(f'Cronograma salvo em: {output_path}')
print(f'Tamanho: {os.path.getsize(output_path)} bytes')
